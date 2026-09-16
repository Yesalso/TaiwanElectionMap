# -*- coding: utf-8 -*-
"""整理臺南市.xlsx：重建「各里彙總」，納入所有候選人（含無黨籍），
並產生「2010新北_得票率.xlsx」格式之得票率地圖檔。"""
import os
import re

from openpyxl import load_workbook
from openpyxl.styles import Font

BASE = r"D:\Windows\TaiwanElection\2024DistrictLegislator\data"
SRC = os.path.join(BASE, "臺南市.xlsx")
OUT_ORG = os.path.join(BASE, "臺南市_整理.xlsx")
OUT_RATE = os.path.join(BASE, "臺南市_得票率.xlsx")

DISTRICTS = ["臺南市第1選舉區", "臺南市第2選舉區", "臺南市第3選舉區",
             "臺南市第4選舉區", "臺南市第5選舉區", "臺南市第6選舉區"]

VALID_HEADER = "有效票數A A=1+2+...+N"

CAND_RE = re.compile(r"^\(\d+\)\s*(.+)$")


def parse_cand_header(h):
    """'(1) 魏耀乾 無' -> ('魏耀乾', '無')"""
    s = CAND_RE.match(str(h)).group(1).strip()
    name, party = s.rsplit(" ", 1)
    return name, party


def read_district(ws):
    hdr = [c.value for c in ws[1]]
    cands = []
    valid_idx = hdr.index(VALID_HEADER)
    for i, h in enumerate(hdr):
        if h and CAND_RE.match(str(h)):
            name, party = parse_cand_header(h)
            cands.append((name, party, i))
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        dist = {name: r[i] for name, party, i in cands}
        rows.append({
            "town": r[0],              # 鄉(鎮、市、區)別
            "village": r[1],           # 村里別
            "votes": dist,
            "valid": r[valid_idx],
            "tail": list(r[valid_idx + 1:]),   # 無效票B ... 投票率H
        })
    return cands, rows


def main():
    wb = load_workbook(SRC)
    all_cands = []          # (district, [(name, party, idx)])
    rows_by_dist = {}

    for d in DISTRICTS:
        cands, rows = read_district(wb[d])
        all_cands.append((d, cands))
        rows_by_dist[d] = rows

    cand_list = [(name, party) for _d, cands in all_cands for name, party, _ in cands]

    # ================= 1) 整理檔：各里彙總重建 =================
    tail_headers = ["無效票數B", "投票數C C=A+B", "已領未投票數 D D=E-C",
                    "發出票數E E=C+D", "用餘票數F", "選舉人數G G=E+F",
                    "投票率H H=C÷G"]
    headers = ["選舉區別", "鄉(鎮、市、區)別", "村里別"] + \
              [f"{n}({p})得票數" for n, p in cand_list] + \
              [VALID_HEADER] + tail_headers

    ws = wb["各里彙總"]
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in row:
            c.value = None

    bold = Font(name="Microsoft JhengHei", size=11, bold=True)
    normal = Font(name="Microsoft JhengHei", size=11)
    for ci, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = bold

    r_i = 2
    for d in DISTRICTS:
        # 該選區候選人對應之欄位（headers 中欄位順序依 cand_list）
        cand_idx = {name: 4 + ci for ci, (name, _p) in enumerate(cand_list)}
        for rec in rows_by_dist[d]:
            ws.cell(row=r_i, column=1, value=d)
            ws.cell(row=r_i, column=2, value=rec["town"])
            ws.cell(row=r_i, column=3, value=rec["village"])
            for name in rec["votes"]:
                ws.cell(row=r_i, column=cand_idx[name], value=rec["votes"][name]).font = normal
            ws.cell(row=r_i, column=4 + len(cand_list), value=rec["valid"])
            for ti, tv in enumerate(rec["tail"], start=5 + len(cand_list)):
                ws.cell(row=r_i, column=ti, value=tv)
            r_i += 1
    wb.save(OUT_ORG)
    print("已產生整理檔:", OUT_ORG, " 候選人欄位數=", len(cand_list), " 里數=", r_i - 2)

    # ================= 2) 得票率地圖檔（仿 2010新北_得票率） =================
    from openpyxl import Workbook
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = "各里彙總"
    rate_headers = ["選舉區別", "鄉(鎮、市、區)別", "村里別"] + \
                   [f"{n}({p})得票率" for n, p in cand_list]
    for ci, h in enumerate(rate_headers, start=1):
        cell = ws2.cell(row=1, column=ci, value=h)
        cell.font = bold

    r_i = 2
    for d in DISTRICTS:
        cand_idx = {name: 4 + ci for ci, (name, _p) in enumerate(cand_list)}
        for rec in rows_by_dist[d]:
            ws2.cell(row=r_i, column=1, value=d).font = normal
            ws2.cell(row=r_i, column=2, value=rec["town"]).font = normal
            ws2.cell(row=r_i, column=3, value=rec["village"]).font = normal
            for name in rec["votes"]:
                v = rec["votes"][name]
                pct = round(v / rec["valid"] * 100, 2) if rec["valid"] else None
                ws2.cell(row=r_i, column=cand_idx[name], value=pct).font = normal
            r_i += 1
    wb2.save(OUT_RATE)
    print("已產生得票率地圖檔:", OUT_RATE, " 里數=", r_i - 2)


if __name__ == "__main__":
    main()