# -*- coding: utf-8 -*-
"""
將「第16任總統副總統選舉候選人在新北市各村(里)得票數一覽表.xlsx」
轉換成與「2014新北_得票率.xlsx」相同的得票率格式。

輸出工作表：各里彙總
欄位：選舉區別、鄉(鎮、市、區)別、村里別、中國國民黨得票率、民主進步黨得票率、台灣民眾黨得票率

得票率(%) = 候選人得票數 ÷ 有效票數A × 100，四捨五入至 2 位小數。

執行方式：
    py convert_2024_president_percent.py
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Font

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(os.path.dirname(BASE_DIR), "data")

CITY = "新北市"

# 政黨欄位順序（與高雄2024.xlsx 各里彙總一致：國民黨 → 民進黨 → 民眾黨）
PARTY_ORDER = ["中國國民黨", "民主進步黨", "台灣民眾黨"]

SOURCE_FILE = os.path.join(
    DATA_DIR,
    "第16任總統副總統選舉候選人在新北市各村(里)得票數一覽表.xlsx",
)
OUTPUT_FILE = os.path.join(DATA_DIR, "2024新北_得票率.xlsx")

VALID_VOTES_COL = 6   # F = 有效票數A


def convert(source_file, output_file):
    import openpyxl
    import warnings
    warnings.filterwarnings("ignore")
    wb = openpyxl.load_workbook(source_file)
    ws = wb["新北市"]

    # 候選人得票欄位於原始資料中的欄位：C=(柯文哲) D=(賴清德) E=(侯友宜)
    CAND_COLS = {
        "台灣民眾黨": 3,
        "民主進步黨": 4,
        "中國國民黨": 5,
    }

    rows = []
    cur_district = None
    for r in range(7, ws.max_row + 1):
        district = ws.cell(r, 1).value
        village = ws.cell(r, 2).value
        if village in (None, ""):
            if district in (None, ""):
                continue
            cur_district = str(district).strip().replace("\u3000", "")
            continue
        valid = ws.cell(r, VALID_VOTES_COL).value
        assert valid in (None, "") or float(valid) > 0
        pcts = []
        for party in PARTY_ORDER:
            votes = float(ws.cell(r, CAND_COLS[party]).value)
            pcts.append(round(votes / float(valid) * 100, 2))
        rows.append([CITY, cur_district, str(village).strip()] + pcts)

    out = Workbook()
    out.remove(out.active)
    ws_out = out.create_sheet("各里彙總")
    headers = ["選舉區別", "鄉(鎮、市、區)別", "村里別"] + [f"{p}得票率" for p in PARTY_ORDER]
    bold = Font(name="Calibri", size=11, bold=True)
    normal = Font(name="Calibri", size=11)
    for c, h in enumerate(headers, start=1):
        cell = ws_out.cell(row=1, column=c, value=h)
        cell.font = bold
    for r, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            cell = ws_out.cell(row=r, column=c, value=v)
            cell.font = normal
    out.save(output_file)
    print(f"已產生: {output_file}  政黨欄位={PARTY_ORDER}  村里數={len(rows)}")


if __name__ == "__main__":
    convert(SOURCE_FILE, OUTPUT_FILE)