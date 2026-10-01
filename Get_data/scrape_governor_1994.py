#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
1994 年台灣省省長選舉 — 縣市 / 鄉鎮市區 / 村里 三層級交叉表爬蟲
=====================================================================
入口：vote3.asp 全國概況（候選人清單，共 5 位）
下鑽：vote31.asp 鄉鎮市區（每頁 333 個鄉鎮市區）→ vote32.asp 村里（每頁約 87 個村里）

1994 年資料庫沒有縣市層頁面（vote3 直接連到 vote31），縣市名需由鄉鎮市區名稱前綴推得：
    臺北縣板橋市 → 縣市 = 臺北縣、鄉鎮市區 = 板橋市
    臺南市北區   → 縣市 = 臺南市、鄉鎮市區 = 北區
縣市清單為 1994 年台灣省轄之 21 個縣、市。

輸出一個 xlsx，含三個工作表，每候選人兩欄（得票數、得票率）：
  - 「縣市級別」    ：縣市 | {候選人}得票數 | {候選人}得票率 | ...
  - 「鄉鎮市區級別」：縣市 | 鄉鎮市區 | {候選人}得票數 | {候選人}得票率 | ...
  - 「村里級別」    ：縣市 | 鄉鎮市區 | 村里 | {候選人}得票數 | {候選人}得票率 | ...
各表末列為「總計」（票數加總；得票率為占全體候選人合計之比率）。

用法：
    python scrape_governor_1994.py [-o 輸出檔名] [--workers 執行緒數] [--delay 秒] [--retry 次]
"""

import argparse
import os
import sys
import threading
import time
from collections import OrderedDict
from concurrent.futures import ThreadPoolExecutor

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

START_URL = "https://vote.nccu.edu.tw/cec/vote3.asp?pass1=L9AA%3CI%3E:88888888iii(("
BASE_URL = "https://vote.nccu.edu.tw/cec/"
OUTPUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "1994台灣省省長選舉_縣市鄉鎮村里.xlsx")

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    )
}

# 1994 年台灣省轄縣市（用於由鄉鎮市區名稱推回縣市）
COUNTIES = [
    "臺北縣", "臺中縣", "臺南縣", "高雄縣", "桃園縣", "苗栗縣", "彰化縣", "南投縣",
    "雲林縣", "嘉義縣", "新竹縣", "屏東縣", "宜蘭縣", "花蓮縣", "臺東縣", "澎湖縣",
    "基隆市", "新竹市", "嘉義市", "臺中市", "臺南市",
]

DELAY = 0.1        # 全域請求間隔（秒），避免對站台造成壓力
WORKERS = 6        # 執行緒數
RETRY = 4

_state = threading.local()
_throttle_lock = threading.Lock()
_next_slot = 0.0


# ---------------------------------------------------------------- 抓取


def get_session():
    """每個執行緒各自持有一個 Session（requests.Session 非執行緒安全）。"""
    s = getattr(_state, "session", None)
    if s is None:
        s = requests.Session()
        s.headers.update(HEADERS)
        _state.session = s
    return s


def throttle():
    """跨執行緒的全域節流：確保相鄰兩次請求至少相隔 DELAY 秒。"""
    global _next_slot
    with _throttle_lock:
        now = time.monotonic()
        wait = _next_slot - now
        _next_slot = max(now, _next_slot) + DELAY
    if wait > 0:
        time.sleep(wait)


def fetch_page(url, retry=RETRY):
    """抓取網頁並以 big5 解碼，失敗自動重試；失敗回傳 None。"""
    for attempt in range(1, retry + 1):
        throttle()
        try:
            resp = get_session().get(url, timeout=30)
            resp.raise_for_status()
            resp.encoding = "big5"
            return resp.text
        except Exception as e:
            print(f"  請求失敗（第 {attempt}/{retry} 次）：{url} -> {type(e).__name__}: {e}")
            if attempt < retry:
                time.sleep(0.8 * attempt)
    return None


def join_url(href):
    if not href:
        return None
    return href if href.startswith("http") else BASE_URL + href


# ---------------------------------------------------------------- 解析


def parse_candidates(html):
    """vote3.asp：候選人清單（10 欄：地區/姓名/號次/性別/出生年次/推薦政黨/得票數/得票率/當選否/是否現任）。"""
    soup = BeautifulSoup(html, "html.parser")
    candidates = []
    seen = set()
    for tr in soup.find_all("tr"):
        cells = tr.find_all("td", recursive=False)
        if len(cells) != 10:
            continue
        a = cells[1].find("a")
        if not a:
            continue
        number = cells[2].get_text(strip=True)
        if number in seen:
            continue
        seen.add(number)
        candidates.append(
            {
                "region": cells[0].get_text(strip=True),
                "name": a.get_text(strip=True),
                "number": number,
                "gender": cells[3].get_text(strip=True),
                "birth": cells[4].get_text(strip=True),
                "party": cells[5].get_text(strip=True),
                "votes": cells[6].get_text(strip=True),
                "rate": cells[7].get_text(strip=True),
                "elected": cells[8].get_text(strip=True),
                "incumbent": cells[9].get_text(strip=True),
                "detail_url": join_url(a.get("href")),
            }
        )
    return candidates


def parse_area_rows(html):
    """解析 vote31 / vote32（5 欄：地區/姓名/號次/得票數/得票率；地區為上層地區相連字串）。"""
    soup = BeautifulSoup(html, "html.parser")
    rows = []
    for tr in soup.find_all("tr"):
        cells = tr.find_all("td", recursive=False)
        vals = [c.get_text(strip=True) for c in cells]
        if len(vals) != 5:
            continue
        area, votes = vals[0], vals[3]
        if not area or not votes.isdigit():
            continue
        a = cells[1].find("a")
        rows.append(
            {
                "area": area,
                "votes": int(votes),
                "rate": parse_rate(vals[4]),
                "detail_url": join_url(a.get("href")) if a else None,
            }
        )
    return rows


def parse_rate(text):
    """'4.31%' → 4.31（四捨五入小數兩位）。"""
    try:
        return round(float(text.rstrip("%")), 2)
    except (ValueError, AttributeError):
        return None


def split_county(area):
    """由鄉鎮市區全名拆出 (縣市, 鄉鎮市區)，例：'臺北縣板橋市' → ('臺北縣', '板橋市')。"""
    for county in COUNTIES:
        if area.startswith(county):
            return county, area[len(county):]
    return "", area


# ---------------------------------------------------------------- 主流程


def scrape(start_url, workers, retry):
    print("抓取入口頁 vote3.asp ...")
    html = fetch_page(start_url, retry=retry)
    if not html:
        print("入口頁抓取失敗，結束。")
        return None

    candidates = parse_candidates(html)
    if not candidates:
        print("未解析到候選人，請確認網址。")
        return None

    print(f"候選人：{len(candidates)} 位")
    for c in candidates:
        print(f"  {c['number']} 號 {c['name']}（{c['party']}）{c['votes']} 票 {c['rate']}")

    county_rates = OrderedDict()      # 縣市 -> None
    town_rates = OrderedDict()        # (縣市, 鄉鎮市區) -> None
    village_rates = OrderedDict()     # (縣市, 鄉鎮市區, 村里) -> None
    county_votes = {}                 # (候選人idx, 縣市) -> (得票數, 得票率)
    town_votes = {}                   # (候選人idx, 縣市, 鄉鎮市區) -> (得票數, 得票率)
    village_votes = {}                # (候選人idx, 縣市, 鄉鎮市區, 村里) -> (得票數, 得票率)

    jobs = []  # (候選人idx, 縣市, 鄉鎮市區, vote32 網址)
    for i, c in enumerate(candidates):
        print(f"\n[{i + 1}/{len(candidates)}] 鄉鎮市區層：{c['name']}（{c['number']}）")
        t_html = fetch_page(c["detail_url"], retry=retry)
        if not t_html:
            continue
        rows = parse_area_rows(t_html)
        print(f"  鄉鎮市區：{len(rows)} 個")
        for row in rows:
            county, township = split_county(row["area"])
            county_rates.setdefault(county, None)
            town_rates.setdefault((county, township), None)
            town_votes[(i, county, township)] = (row["votes"], row["rate"])
            if row["detail_url"]:
                jobs.append((i, county, township, row["detail_url"]))

    # ---- 村里層：多執行緒並行抓取 ----
    print(f"\n抓取村里層：{len(jobs)} 個鄉鎮市區，{workers} 執行緒 ...")
    done = [0]
    total = len(jobs)
    failed = [0]
    done_lock = threading.Lock()
    start = time.time()

    def worker(job):
        i, county, township, url = job
        v_html = fetch_page(url, retry=retry)
        if v_html:
            prefix = county + township
            with done_lock:
                for v in parse_area_rows(v_html):
                    village = v["area"][len(prefix):] if v["area"].startswith(prefix) else v["area"]
                    village_rates.setdefault((county, township, village), None)
                    village_votes[(i, county, township, village)] = (v["votes"], v["rate"])
        with done_lock:
            done[0] += 1
            if not v_html:
                failed[0] += 1
            n = done[0]
        if n % 50 == 0 or n == total:
            elapsed = time.time() - start
            speed = n / elapsed if elapsed else 0
            eta = (total - n) / speed if speed else 0
            print(f"  進度 [{n}/{total}]  {elapsed:6.0f} 秒  {speed:4.1f} 頁/秒  剩餘約 {eta / 60:4.1f} 分")

    if jobs:
        with ThreadPoolExecutor(max_workers=workers) as pool:
            list(pool.map(worker, jobs))

    # 由鄉鎮市區加總補上縣市層的得票數與得票率
    for county in county_rates:
        per_cand = {}
        for i in range(len(candidates)):
            per_cand[i] = sum(
                votes for (ci, co, _tw), (votes, _rate) in town_votes.items()
                if ci == i and co == county
            )
        grand = sum(per_cand.values())
        for i, v in per_cand.items():
            county_votes[(i, county)] = (v, round(v / grand * 100, 2) if grand else None)

    if failed[0]:
        print(f"  注意：{failed[0]} 個鄉鎮市區的村里頁抓取失敗。")

    return {
        "candidates": candidates,
        "counties": list(county_rates.keys()),
        "towns": list(town_rates.keys()),
        "villages": list(village_rates.keys()),
        "county_votes": county_votes,
        "town_votes": town_votes,
        "village_votes": village_votes,
    }


# ---------------------------------------------------------------- 輸出


def write_xlsx(data, path):
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    total_font = Font(bold=True)

    candidates = data["candidates"]
    n_cand = len(candidates)

    wb = Workbook()
    wb.remove(wb.active)

    def build_sheet(title, prefixes, parts_fn, keys, votes_lookup, key_fn, widths):
        """每候選人兩欄（得票數、得票率）交錯排列，與其他爬蟲輸出格式一致。"""
        ws = wb.create_sheet(title)
        headers = list(prefixes)
        for c in candidates:
            headers.append(f"{c['name']}得票數")
            headers.append(f"{c['name']}得票率")
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        ws.freeze_panes = ws.cell(row=2, column=len(prefixes) + 1)
        ws.row_dimensions[1].height = 30

        np = len(prefixes)

        def col_of(i, rate):
            return np + 1 + i * 2 + (1 if rate else 0)

        for r, key in enumerate(keys, 2):
            for col, part in enumerate(parts_fn(key), 1):
                ws.cell(row=r, column=col, value=part).border = border
            for i, _c in enumerate(candidates):
                votes, rate = votes_lookup.get(key_fn(i, key), ("", ""))
                ws.cell(row=r, column=col_of(i, False), value=votes).border = border
                ws.cell(row=r, column=col_of(i, True), value=rate).border = border

        # 總計列
        totals = [0] * n_cand
        for key in keys:
            for i in range(n_cand):
                totals[i] += votes_lookup.get(key_fn(i, key), (0, 0))[0] or 0
        grand = sum(totals)

        tr = len(keys) + 2
        ws.cell(row=tr, column=1, value="總計").font = total_font
        for i in range(n_cand):
            v = totals[i]
            ws.cell(row=tr, column=col_of(i, False), value=v).font = total_font
            rate = round(v / grand * 100, 2) if grand else ""
            ws.cell(row=tr, column=col_of(i, True), value=rate).font = total_font

        for col, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = w
        for i in range(n_cand):
            ws.column_dimensions[get_column_letter(col_of(i, False))].width = 12
            ws.column_dimensions[get_column_letter(col_of(i, True))].width = 10

        print(f"  「{title}」：{len(keys)} 列 × {n_cand} 位候選人")

    build_sheet(
        "縣市級別", ["縣市"], lambda k: [k], data["counties"], data["county_votes"],
        lambda i, k: (i, k), [16],
    )
    build_sheet(
        "鄉鎮市區級別", ["縣市", "鄉鎮市區"], lambda k: [k[0], k[1]], data["towns"], data["town_votes"],
        lambda i, k: (i, k[0], k[1]), [14, 14],
    )
    build_sheet(
        "村里級別", ["縣市", "鄉鎮市區", "村里別"], lambda k: [k[0], k[1], k[2]],
        data["villages"], data["village_votes"],
        lambda i, k: (i, k[0], k[1], k[2]), [14, 14, 14],
    )

    wb.save(path)
    print(f"\n完成！已儲存至 '{path}'")


def verify(data):
    """以各層加總票數與 vote3 全國公告票數核對。"""
    print("\n票數核對（鄉鎮市區加總 vs 全國公告）：")
    votes = [0] * len(data["candidates"])
    for (i, _co, _tw), (v, _r) in data["town_votes"].items():
        votes[i] += v or 0
    for c, v in zip(data["candidates"], votes):
        try:
            official = int(str(c["votes"]).replace(",", ""))
        except ValueError:
            official = None
        if official is None:
            print(f"  {c['name']}（{c['number']}）：公告票數無法解析")
            continue
        mark = "一致" if official == v else f"不一致（差 {v - official:+d}）"
        print(f"  {c['name']}（{c['number']}）：加總 {v} vs 公告 {official} → {mark}")


# ---------------------------------------------------------------- CLI


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description="爬取 1994 年台灣省省長選舉 縣市/鄉鎮市區/村里 三層級交叉表")
    parser.add_argument("start_url", nargs="?", default=START_URL, help="vote3.asp 入口網址")
    parser.add_argument("-o", "--output", default=OUTPUT, help=f"輸出檔名（預設 {os.path.basename(OUTPUT)}）")
    parser.add_argument("--workers", type=int, default=WORKERS, help=f"並行執行緒數（預設 {WORKERS}）")
    parser.add_argument("--delay", type=float, default=DELAY, help=f"全域請求間隔秒數（預設 {DELAY}）")
    parser.add_argument("--retry", type=int, default=RETRY, help=f"失敗重試次數（預設 {RETRY}）")
    return parser.parse_args(argv)


def main(argv=None):
    global DELAY
    args = parse_args(argv)
    DELAY = args.delay

    output = args.output if os.path.isabs(args.output) else os.path.join(
        os.path.dirname(os.path.abspath(__file__)), args.output
    )
    if os.path.exists(output):
        try:
            os.remove(output)
        except PermissionError:
            base, ext = os.path.splitext(output)
            output = f"{base}_新{ext}"
            print(f"原檔被佔用，改存為 '{output}'")

    start = time.time()
    data = scrape(args.start_url, args.workers, args.retry)
    if not data:
        return 1
    verify(data)
    write_xlsx(data, output)
    print(f"  總耗時：{time.time() - start:.0f} 秒")
    return 0


if __name__ == "__main__":
    sys.exit(main())
