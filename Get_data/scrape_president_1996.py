#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Scrape the 1996 ROC presidential election from NCCU CEC
and output a flat xlsx with 10 columns:
  地區, 姓名, 號次, 性別, 出生年次, 推薦政黨, 得票數, 得票率, 當選否, 是否現任
"""

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

URL = r"https://vote.nccu.edu.tw/cec/vote3.asp?pass1=I9AA%3EI8888888888iii(("
OUTPUT = r"D:\Windows\TaiwanElection\Get_data\1996總統副總統選舉.xlsx"

HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"}

COL_NAMES = ["地區", "姓名", "號次", "性別", "出生年次", "推薦政黨", "得票數", "得票率", "當選否", "是否現任"]


def fetch(url):
    r = requests.get(url, headers=HEADERS, timeout=30)
    r.encoding = "big5"
    return r.text


def parse(html):
    soup = BeautifulSoup(html, "html.parser")
    rows = soup.find("table").find_all("tr")
    data = []
    pending = None  # carries over party/votes/rate/elected/incumbent from a 10-cell row
    for tr in rows[1:]:  # skip header
        cells = tr.find_all("td", recursive=False)
        vals = [c.get_text(strip=True) for c in cells]
        if len(vals) == 10:
            # President row: 地區,姓名,號次,性別,出生年次,政黨,得票數,得票率,當選否,是否現任
            data.append(vals)
            # Store region + shared fields for the VP row that follows
            pending = [vals[0]] + vals[5:]  # [地區,政黨,得票數,得票率,當選否,是否現任]
        elif len(vals) == 4 and pending:
            # VP row: 姓名,號次,性別,出生年次
            data.append(pending[:1] + vals[:1] + vals[1:4] + pending[1:])
            pending = None
    return data


def write_xlsx(data, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "總統副總統得票"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    # header
    for c, name in enumerate(COL_NAMES, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    # data
    for r, row in enumerate(data, 2):
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = border
            if c == 1:
                pass  # region left-aligned

    # column widths
    widths = [14, 10, 6, 6, 10, 28, 12, 10, 8, 8]
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    wb.save(path)
    print(f"Saved: {path}")


def main():
    print("Fetching...")
    html = fetch(URL)
    data = parse(html)
    print(f"Parsed {len(data)} rows")
    for row in data:
        print(row)
    write_xlsx(data, OUTPUT)


if __name__ == "__main__":
    main()
