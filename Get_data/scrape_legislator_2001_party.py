#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
第 05 屆立法委員選舉（2001）— 政黨 × 鄉鎮市區 得票地圖
====================================================
入口：vote3.asp（候選人得票明細）
下鑽：每一位候選人 -> vote31.asp（鄉鎮市區層級得票）

因為同一政黨在同一選區可能提名多位候選人，本程式將同一政黨在「同一個鄉鎮市區」
的所有候選人得票加總，輸出「政黨 × 鄉鎮市區」的得票交叉表（每政黨兩欄：得票數、得票率）。

輸出 xlsx，含兩個工作表：
  - 鄉鎮市區級別：縣市 | 鄉鎮市區 | {政黨}得票數/得票率 ... | 總得票數
  - 縣市級別：縣市 | {政黨}得票數/得票率 ... | 總得票數

用法：
    python scrape_legislator_2001_party.py
"""

import os
import time
from collections import OrderedDict

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

START_URL = "https://vote.nccu.edu.tw/cec/vote3.asp?pass1=J:889I8888888888iii(("
BASE_URL = "https://vote.nccu.edu.tw/cec/"
OUTPUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "2001立法委員_政黨鄉鎮得票.xlsx")

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    )
}
DELAY = 0.2
RETRY = 4

COUNTIES = [
    "臺北市", "高雄市", "臺北縣", "宜蘭縣", "桃園縣", "新竹縣", "苗栗縣", "臺中縣",
    "彰化縣", "南投縣", "雲林縣", "嘉義縣", "臺南縣", "高雄縣", "屏東縣", "臺東縣",
    "花蓮縣", "澎湖縣", "基隆市", "新竹市", "臺中市", "嘉義市", "臺南市", "金門縣", "連江縣",
]


def fetch_page(session, url, retry=RETRY):
    for attempt in range(1, retry + 1):
        try:
            r = session.get(url, headers=HEADERS, timeout=30)
            r.raise_for_status()
            r.encoding = "big5"
            return r.text
        except Exception as e:
            print(f"  請求失敗（{attempt}/{retry}）：{url} -> {e}")
            if attempt < retry:
                time.sleep(DELAY * 3)
    return None


def join_url(href):
    if not href:
        return None
    return href if href.startswith("http") else BASE_URL + href


def county_of(district):
    for c in COUNTIES:
        if district.startswith(c):
            return c
    return district


def parse_candidates(html):
    """vote3.asp：所有候選人（行政區/姓名/號次/政黨/得票數/得票率 + vote31 網址）"""
    soup = BeautifulSoup(html, "html.parser")
    cands = []
    for tr in soup.find("table").find_all("tr"):
        cells = tr.find_all("td", recursive=False)
        vals = [c.get_text(strip=True) for c in cells]
        if len(vals) < 8:
            continue
        a = cells[1].find("a")
        if not a:
            continue
        cands.append(
            {
                "district": vals[0],
                "name": a.get_text(strip=True),
                "number": vals[2],
                "party": vals[5],
                "votes": vals[6],
                "rate": vals[7],
                "detail_url": join_url(a.get("href")),
            }
        )
    return cands


def parse_township_rows(html, district):
    """vote31.asp：回傳 [(鄉鎮市區, 得票數)]"""
    soup = BeautifulSoup(html, "html.parser")
    table = soup.find("table")
    rows = []
    if not table:
        return rows
    for tr in table.find_all("tr")[1:]:  # skip header
        cells = tr.find_all("td", recursive=False)
        vals = [c.get_text(strip=True) for c in cells]
        if len(vals) < 4:
            continue
        area = vals[0]
        votes = vals[3]
        if not area or not votes.isdigit():
            continue
        township = area[len(district):]
        rows.append((township, int(votes)))
    return rows


def write_xlsx(county_totals, town_rows, parties, path):
    wb = Workbook()
    wb.remove(wb.active)

    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    def style_header(ws, headers):
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center",
                                       wrap_text=True)
            cell.border = border

    def fill_body(ws, prefix, data_keys, data_lookup):
        pcol = len(prefix(data_keys[0]))  # 前置欄數
        for r, key in enumerate(data_keys, 2):
            d = data_lookup.get(key) or {}
            total = d.get("_TOTAL_", 0)
            for pc, val in enumerate(prefix(key), 1):
                cell = ws.cell(row=r, column=pc, value=val)
                cell.border = border
            for pi, p in enumerate(parties):
                votes = d.get(p, 0)
                cell = ws.cell(row=r, column=pcol + 1 + pi * 2, value=votes)
                cell.border = border
                rate = f"{votes / total * 100:.2f}%" if total else "-"
                cell2 = ws.cell(row=r, column=pcol + 2 + pi * 2, value=rate)
                cell2.border = border
            ws.cell(row=r, column=pcol + 1 + len(parties) * 2, value=total).border = border

    def widths(ws, prefix_len, parties):
        for i in range(prefix_len):
            ws.column_dimensions[get_column_letter(i + 1)].width = 14
        last = prefix_len + len(parties) * 2 + 1
        for i in range(len(parties)):
            ws.column_dimensions[get_column_letter(prefix_len + 1 + i * 2 + 1)].width = 12
            ws.column_dimensions[get_column_letter(prefix_len + 1 + i * 2 + 2)].width = 10
        ws.column_dimensions[get_column_letter(last)].width = 10

    # ---------- Sheet 1：鄉鎮市區級別 ----------
    ws1 = wb.create_sheet("鄉鎮市區級別")
    headers1 = ["縣市", "鄉鎮市區"]
    for p in parties:
        headers1.append(f"{p}得票數")
        headers1.append(f"{p}得票率")
    headers1.append("總得票數")
    style_header(ws1, headers1)

    town_keys = sorted(town_rows.keys())
    fill_body(ws1, lambda k: [k[0], k[1]], town_keys, town_rows)
    widths(ws1, 2, parties)
    ws1.freeze_panes = "C2"

    # ---------- Sheet 2：縣市級別 ----------
    ws2 = wb.create_sheet("縣市級別")
    headers2 = ["縣市"]
    for p in parties:
        headers2.append(f"{p}得票數")
        headers2.append(f"{p}得票率")
    headers2.append("總得票數")
    style_header(ws2, headers2)

    county_keys = sorted(county_totals.keys())
    fill_body(ws2, lambda k: [k], county_keys, county_totals)
    widths(ws2, 1, parties)
    ws2.freeze_panes = "B2"

    wb.save(path)
    print(f"\n完成！已儲存至 '{path}'")
    print(f"  政黨：{len(parties)} 個")
    print(f"  鄉鎮市區級別：{len(town_keys)} 個鄉鎮市區")
    print(f"  縣市級別：{len(county_keys)} 個縣市")


def main():
    session = requests.Session()
    print("抓取入口頁 vote3.asp ...")
    html = fetch_page(session, START_URL)
    if not html:
        print("入口頁抓取失敗，結束。")
        return 1

    cands = parse_candidates(html)
    print(f"候選人：{len(cands)} 位")

    parties = list(OrderedDict.fromkeys(c["party"] for c in cands))
    print(f"政黨：{len(parties)} 個")
    for p in parties:
        n = sum(1 for c in cands if c["party"] == p)
        print(f"  {p}: {n} 人")

    town_rows = {}  # (縣市, 鄉鎮市區) -> {政黨: 得票數, _TOTAL_: 總數}
    county_totals = {}  # 縣市 -> {政黨: 得票數, _TOTAL_: 總數}
    failed = []

    for i, c in enumerate(cands, 1):
        url = c["detail_url"]
        if not url:
            continue
        html = fetch_page(session, url)
        if not html:
            failed.append(c["name"])
            continue
        co = county_of(c["district"])
        for township, votes in parse_township_rows(html, c["district"]):
            key = (co, township)
            d = town_rows.setdefault(key, {})
            d[c["party"]] = d.get(c["party"], 0) + votes
            d["_TOTAL_"] = d.get("_TOTAL_", 0) + votes
            cd = county_totals.setdefault(co, {})
            cd[c["party"]] = cd.get(c["party"], 0) + votes
            cd["_TOTAL_"] = cd.get("_TOTAL_", 0) + votes
        time.sleep(DELAY)
        if i % 50 == 0 or i == len(cands):
            print(f"  進度：{i}/{len(cands)}")

    if failed:
        print("抓取失敗的候選人：", failed)

    write_xlsx(county_totals, town_rows, parties, OUTPUT)
    return 0


if __name__ == "__main__":
    import sys
    sys.exit(main())