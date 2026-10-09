#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
第 14 任總統副總統選舉（2016）— 縣市 / 鄉鎮市區 / 村里 三層級得票率爬蟲
===============================================================================
入口：vote3.asp 全國概況（候選人清單）
下鑽：vote312.asp 縣市 → vote313.asp 鄉鎮市區 → vote32.asp 村里

輸出一個 xlsx，含三個工作表，每候選人一欄（僅得票率，單位為 %）：
  - 「縣市級別」：選舉區別 | {候選人}得票率 ...
  - 「鄉鎮市區級別」：選舉區別 | 鄉(鎮、市、區)別 | {候選人}得票率 ...
  - 「村里級別」：選舉區別 | 鄉(鎮、市、區)別 | 村里別 | {候選人}得票率 ...

用法：
    python scrape_president_2016_rates.py [-o 輸出檔名] [--delay 秒]
"""

import argparse
import os
import sys
import time
from collections import OrderedDict

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

START_URL = "https://vote.nccu.edu.tw/cec/vote3.asp?pass1=I:89%3EI8888888888iii(("
BASE_URL = "https://vote.nccu.edu.tw/cec/"
OUTPUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Precident", "2016Precident", "data")
OUTPUT = os.path.join(OUTPUT_DIR, "2016總統副總統選舉_得票率.xlsx")

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    )
}
DELAY = 0.2
RETRY = 4


def fetch_page(session, url, retry=RETRY):
    for attempt in range(1, retry + 1):
        try:
            r = session.get(url, headers=HEADERS, timeout=30)
            r.raise_for_status()
            r.encoding = "big5"
            return r.text
        except Exception as e:
            print(f"  請求失敗（{attempt}/{retry}）：{url} -> {e}")
            if attempt < retry:
                time.sleep(DELAY * 3)
    return None


def join_url(href):
    if not href:
        return None
    return href if href.startswith("http") else BASE_URL + href


def parse_candidates(html):
    """vote3.asp：候選人清單（每組正副總統取總統為代表，避免重複）。"""
    soup = BeautifulSoup(html, "html.parser")
    candidates = []
    seen = set()
    for tr in soup.find_all("tr"):
        cells = tr.find_all("td", recursive=False)
        vals = [c.get_text(strip=True) for c in cells]
        if len(vals) != 10:
            continue
        a = cells[1].find("a")
        if not a:
            continue
        number = vals[2]
        if number in seen:
            continue
        seen.add(number)
        candidates.append(
            {
                "name": vals[1],
                "number": number,
                "party": vals[5],
                "detail_url": join_url(a.get("href")),
            }
        )
    return candidates


def parse_area_rows(html):
    """解析 vote312 / vote313 / vote32（5 欄：地區 / 姓名 / 號次 / 得票數 / 得票率）。"""
    soup = BeautifulSoup(html, "html.parser")
    rows = []
    for tr in soup.find_all("tr"):
        cells = tr.find_all("td", recursive=False)
        vals = [c.get_text(strip=True) for c in cells]
        if len(vals) != 5:
            continue
        area, votes = vals[0], vals[3]
        if not area or not votes.isdigit():
            continue
        a = cells[1].find("a")
        rows.append(
            {
                "area": area,
                "votes": int(votes),
                "rate": parse_rate(vals[4]),
                "detail_url": join_url(a.get("href")) if a else None,
            }
        )
    return rows


def parse_rate(text):
    """將 '4.02542372881356%' 轉為四捨五入到小數兩位的 float（4.03）。"""
    try:
        return round(float(text.rstrip("%")), 2)
    except (ValueError, AttributeError):
        return None


def scrape(session, delay):
    print("抓取入口頁 vote3.asp ...")
    html = fetch_page(session, START_URL)
    if not html:
        print("入口頁抓取失敗，結束。")
        return None

    candidates = parse_candidates(html)
    print(f"候選人：{len(candidates)} 組")
    for c in candidates:
        print(f"  {c['number']} 號 {c['name']}（{c['party']}）")

    counties = OrderedDict()
    towns = OrderedDict()
    villages = OrderedDict()
    county_rates = {}
    town_rates = {}
    village_rates = {}

    for i, c in enumerate(candidates):
        print(f"\n[{i + 1}/{len(candidates)}] {c['name']}（{c['number']}）")
        county_html = fetch_page(session, c["detail_url"])
        if not county_html:
            continue
        county_rows = parse_area_rows(county_html)
        print(f"  縣市：{len(county_rows)} 個")

        for county_row in county_rows:
            county = county_row["area"]
            counties.setdefault(county, None)
            county_rates[(i, county)] = county_row["rate"]
            if not county_row["detail_url"]:
                continue

            town_html = fetch_page(session, county_row["detail_url"])
            if not town_html:
                continue
            for town_row in parse_area_rows(town_html):
                township = town_row["area"][len(county):]
                towns.setdefault((county, township), None)
                town_rates[(i, county, township)] = town_row["rate"]
                if not town_row["detail_url"]:
                    continue

                village_html = fetch_page(session, town_row["detail_url"])
                if not village_html:
                    continue
                prefix = county + township
                for village_row in parse_area_rows(village_html):
                    village = village_row["area"][len(prefix):]
                    villages.setdefault((county, township, village), None)
                    village_rates[(i, county, township, village)] = village_row["rate"]
                time.sleep(delay)

            print(f"    {county}：完成")
            time.sleep(delay)

        time.sleep(delay)

    return {
        "candidates": candidates,
        "counties": list(counties.keys()),
        "towns": list(towns.keys()),
        "villages": list(villages.keys()),
        "county_rates": county_rates,
        "town_rates": town_rates,
        "village_rates": village_rates,
    }


def write_xlsx(data, path):
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    wb = Workbook()
    wb.remove(wb.active)
    candidates = data["candidates"]

    def style_header(ws, headers):
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        ws.freeze_panes = "A2"

    def rate_headers(prefix):
        headers = list(prefix)
        for c in candidates:
            headers.append(f"{c['name']}得票率")
        return headers

    def fill(ws, keys, prefix_fn, rate_lookup, rate_key_fn, prefix_widths):
        for r, key in enumerate(keys, 2):
            for col, val in enumerate(prefix_fn(key), 1):
                ws.cell(row=r, column=col, value=val).border = border
            for ci, _c in enumerate(candidates):
                val = rate_lookup.get(rate_key_fn(ci, key), "")
                ws.cell(row=r, column=len(prefix_fn(key)) + 1 + ci, value=val).border = border
        for idx, w in enumerate(prefix_widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = w
        for ci in range(len(candidates)):
            ws.column_dimensions[get_column_letter(len(prefix_widths) + 1 + ci)].width = 14

    ws1 = wb.create_sheet("縣市級別")
    style_header(ws1, rate_headers(["選舉區別"]))
    fill(
        ws1,
        data["counties"],
        lambda co: [co],
        data["county_rates"],
        lambda ci, co: (ci, co),
        [16],
    )

    ws2 = wb.create_sheet("鄉鎮市區級別")
    style_header(ws2, rate_headers(["選舉區別", "鄉(鎮、市、區)別"]))
    fill(
        ws2,
        data["towns"],
        lambda k: [k[0], k[1]],
        data["town_rates"],
        lambda ci, k: (ci, k[0], k[1]),
        [14, 14],
    )

    ws3 = wb.create_sheet("村里級別")
    style_header(ws3, rate_headers(["選舉區別", "鄉(鎮、市、區)別", "村里別"]))
    fill(
        ws3,
        data["villages"],
        lambda k: [k[0], k[1], k[2]],
        data["village_rates"],
        lambda ci, k: (ci, k[0], k[1], k[2]),
        [14, 14, 14],
    )

    wb.save(path)
    print(f"\n完成！已儲存至 '{path}'")
    print(f"  縣市級別：{len(data['counties'])} 個")
    print(f"  鄉鎮市區級別：{len(data['towns'])} 個")
    print(f"  村里級別：{len(data['villages'])} 個")


def main(argv=None):
    parser = argparse.ArgumentParser(description="爬取第14任總統副總統選舉（2016）三層級得票率")
    parser.add_argument("-o", "--output", default=OUTPUT, help=f"輸出檔名（預設 {OUTPUT}）")
    parser.add_argument("--delay", type=float, default=DELAY, help=f"每頁間隔秒數（預設 {DELAY}）")
    args = parser.parse_args(argv)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    output = args.output if os.path.isabs(args.output) else os.path.join(OUTPUT_DIR, args.output)

    session = requests.Session()
    start = time.time()
    data = scrape(session, args.delay)
    if not data:
        return 1
    write_xlsx(data, output)
    print(f"  總耗時：{time.time() - start:.0f} 秒")
    return 0


if __name__ == "__main__":
    sys.exit(main())